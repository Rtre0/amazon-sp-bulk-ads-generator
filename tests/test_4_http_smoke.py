# -*- coding: utf-8 -*-
"""v6 HTTP 冒烟测试：/api/analyze -> /api/generate -> /api/dictionary 增删查 / import / export。"""
import json
import os
import sys
import time
import urllib.parse
import urllib.request
import uuid

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
API = "http://127.0.0.1:8000"
SAMPLE = os.path.join(BASE, "samples", "多店铺测试样本", "赛狐导出_多店铺日文品名.xlsx")
sys.stdout.reconfigure(encoding="utf-8")

OK = FAIL = 0


def check(name, cond, detail=""):
    global OK, FAIL
    if cond:
        OK += 1
        print(f"  [PASS] {name} {detail}")
    else:
        FAIL += 1
        print(f"  [FAIL] {name} {detail}")


def post_multipart(url, fields, files):
    boundary = "----wb" + uuid.uuid4().hex
    body = b""
    for k, v in fields.items():
        body += f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n".encode()
    for k, (fn, data) in files.items():
        body += (
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"; "
            f"filename=\"{fn}\"\r\nContent-Type: application/octet-stream\r\n\r\n"
        ).encode() + data + b"\r\n"
    body += f"--{boundary}--\r\n".encode()
    req = urllib.request.Request(
        url, data=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST",
    )
    return urllib.request.urlopen(req, timeout=180)


def post_json(url, obj, method="POST"):
    req = urllib.request.Request(
        url, data=json.dumps(obj).encode("utf-8"),
        headers={"Content-Type": "application/json"}, method=method,
    )
    return urllib.request.urlopen(req, timeout=60)


print("=" * 78)
print("v6 HTTP 冒烟测试")
print("=" * 78)

# ---------------- /api/analyze ----------------
print("\n[1] POST /api/analyze（上传赛狐多店铺样本）")
raw = open(SAMPLE, "rb").read()
t0 = time.time()
resp = post_multipart(f"{API}/api/analyze", {},
                      {"file": ("赛狐导出_多店铺日文品名.xlsx", raw)})
data = json.loads(resp.read().decode("utf-8"))
dt = time.time() - t0
shops = data["shops"]
print(f"     耗时 {dt:.2f}s  表头行 {data['header_row']}  店铺 {list(shops)}")
print(f"     解析列: {data['resolved_columns']}")
check("HTTP 200", resp.status == 200)
check("返回 4 个店铺", len(shops) == 4, str(list(shops)))
check("解析耗时 < 2s（纯离线）", dt < 2.0, f"{dt:.2f}s")
p0 = shops[list(shops)[0]]["products"][0]
check("商品含中文品名", bool(p0.get("product_name")), p0.get("product_name"))
check("商品含真实均价", float(p0.get("avg_price", 0)) > 0, str(p0.get("avg_price")))
check("商品含子体数", int(p0.get("child_count", 0)) > 0, str(p0.get("child_count")))

# ---------------- /api/generate ----------------
print("\n[2] POST /api/generate（生成 Bulk Excel 下载）")
all_products = [p for s in shops.values() for p in s["products"]]
payload = {
    "products": all_products,
    "settings": {"include_targeting": True, "include_placement": True,
                 "placement_top_pct": 30, "placement_product_pct": 10,
                 "default_bid": 25, "bidding_strategy": "Fixed bid"},
}
resp2 = post_json(f"{API}/api/generate", payload)
blob = resp2.read()
out = os.path.join(BASE, "samples", "v6测试样本")
os.makedirs(out, exist_ok=True)
out_path = os.path.join(out, "生成的Bulk_v6_HTTP.xlsx")
open(out_path, "wb").write(blob)
print(f"     下载 {len(blob)} 字节 -> {os.path.relpath(out_path, BASE)}")
check("HTTP 200", resp2.status == 200)
check("Content-Type 为 xlsx",
      "spreadsheetml" in resp2.headers.get("Content-Type", ""),
      resp2.headers.get("Content-Type", ""))
check("文件非空", len(blob) > 10000, f"{len(blob)} 字节")

import openpyxl  # noqa: E402

from core.bulk_generator import OFFICIAL_COLUMNS  # noqa: E402

ws = openpyxl.load_workbook(out_path)["Sponsored Products Campaigns"]
print(f"     行数 {ws.max_row - 1}  列数 {ws.max_column}")
check("列数 32（31 官方列 + ASIN 辅助）", ws.max_column == 32, str(ws.max_column))
check("行数 = 583", ws.max_row - 1 == 583, str(ws.max_row - 1))
_hdr = [c.value for c in ws[1]]
check("表头前 31 列 = 官方模板", _hdr[:31] == OFFICIAL_COLUMNS, str(_hdr[25:28]))
check("第 27 列为 拓展商品投放编号", _hdr[26] == "拓展商品投放编号", repr(_hdr[26]))

# ---------------- 词库接口 ----------------
print("\n[3] 词库统计与单条增删改")
before = json.loads(urllib.request.urlopen(f"{API}/api/dictionary/stats", timeout=20).read())["total_count"]
check("GET /api/dictionary/stats", before > 0, f"{before} 条")
full = json.loads(urllib.request.urlopen(f"{API}/api/dictionary", timeout=20).read())
check("GET /api/dictionary 返回 entries", len(full.get("entries", [])) == before, f"{len(full.get('entries', []))}")
check("GET /api/dictionary 返回黑名单规模", full.get("blacklist_size", 0) > 0, str(full.get("blacklist_size")))

r = json.loads(post_json(f"{API}/api/dictionary", {"ja": "テスト名詞", "zh": "测试名词"}).read())
check("POST 新增成功", r.get("ok") is True and r["total_count"] == before + 1,
      f"{before} -> {r['total_count']}")

# 新增后应能立即驱动解析
resp3 = post_multipart(f"{API}/api/analyze", {}, {"file": ("s.xlsx", raw)})
d3 = json.loads(resp3.read().decode("utf-8"))
check("新增词条后解析仍正常", len(d3["shops"]) == 4, str(list(d3["shops"])))

req = urllib.request.Request(
    f"{API}/api/dictionary?ja=" + urllib.parse.quote("テスト名詞"), method="DELETE")
rd = json.loads(urllib.request.urlopen(req, timeout=20).read())
check("DELETE 成功", rd.get("ok") is True and rd["total_count"] == before, str(rd["total_count"]))
rr = json.loads(post_json(f"{API}/api/dictionary/reload", {}).read())
check("POST /api/dictionary/reload", rr.get("ok") is True and rr["total_count"] == before, str(rr["total_count"]))

# ---------------- 一键导入 ----------------
print("\n[4] POST /api/dictionary/import（一键导入对照表 JSON）")
# 用随机后缀的唯一键，保证 imported_count 可预期（不受既有词库状态影响）
suffix = uuid.uuid4().hex[:6]
k1, k2 = f"テスト陳列棚{suffix}", f"テスト包丁立て{suffix}"
import_payload = {
    k1: "测试陈列架",
    k2: "测试菜刀架",
    "折りたたみ": "折叠",     # 应被黑名单拒绝
    "大型": "大型",           # 应被黑名单拒绝
}
resp4 = post_multipart(f"{API}/api/dictionary/import", {},
                       {"file": ("custom_nouns.json", json.dumps(import_payload, ensure_ascii=False).encode("utf-8"))})
imp = json.loads(resp4.read().decode("utf-8"))
print(f"     {imp}")
check("HTTP 200", resp4.status == 200)
check("status = success", imp.get("status") == "success", str(imp.get("status")))
check("导入计数 = 2", imp.get("imported_count") == 2, str(imp.get("imported_count")))
check("黑名单拒绝 2 条", sorted(imp.get("purged_keys", [])) == sorted(["折りたたみ", "大型"]),
      str(sorted(imp.get("purged_keys", []))))
check("total_count 正确", imp.get("total_count") == before + 2, str(imp.get("total_count")))

# 导入后应立即可用于解析（新词条能被命中）
resp5 = post_multipart(f"{API}/api/analyze", {}, {"file": ("s.xlsx", raw)})
d5 = json.loads(resp5.read().decode("utf-8"))
names = {p["parent_sku"]: p["product_name"] for s in d5["shops"].values() for p in s["products"]}
print(f"     ff061 -> {names.get('ff061')}")
check("解析仍正常", names.get("ff061") == "加湿器", str(names.get("ff061")))

# 清理导入的测试词条
for k in (k1, k2):
    urllib.request.urlopen(urllib.request.Request(
        f"{API}/api/dictionary?ja=" + urllib.parse.quote(k), method="DELETE"), timeout=20).read()
after = json.loads(urllib.request.urlopen(f"{API}/api/dictionary/stats", timeout=20).read())["total_count"]
check("清理后恢复原条目数", after == before, f"{before} -> {after}")

# ---------------- 导出 ----------------
print("\n[5] GET /api/dictionary/export")
exp = json.loads(urllib.request.urlopen(f"{API}/api/dictionary/export", timeout=20).read())
check("导出条目数与词库一致", len(exp) == after, f"{len(exp)} / {after}")
check("导出内容可读", "包丁立て" in exp and exp["包丁立て"] == "菜刀架", str(exp.get("包丁立て")))

# ---------------- v7 一键刷新品名匹配 ----------------
print("\n[6] POST /api/dictionary/rematch（数量词穿透 + 词库动态刷新）")
rm_payload = {
    "items": [
        {"key": "shopA||mj355", "parent_sku": "mj355",
         "raw_name": "FryWare 2個 マッシャー ステンレス製 サツマイモつぶし"},
        {"key": "shopA||llc16", "parent_sku": "llc16",
         "raw_name": "FryWare ビニール袋 鮮度保持袋 ロール ポリ袋"},
        {"key": "shopB||mj355", "parent_sku": "mj355",
         "raw_name": "FryWare 2個 収納ボックス ジュエリーオーガナイザー"},
        {"parent_sku": "nomatch", "raw_name": "謎の物体xyz"},
        {"parent_sku": "blank", "raw_name": "   "},
    ]
}
resp6 = post_json(f"{API}/api/dictionary/rematch", rm_payload)
rm = json.loads(resp6.read().decode("utf-8"))
print(f"     {rm}")
m = rm.get("matched_names", {})
check("HTTP 200", resp6.status == 200)
check("key 作主键：shopA||mj355 -> 压泥器（跳过 2個）",
      m.get("shopA||mj355") == "压泥器", str(m.get("shopA||mj355")))
check("shopA||llc16 -> 塑料袋", m.get("shopA||llc16") == "塑料袋", str(m.get("shopA||llc16")))
check("跨店铺同名 SKU 不串味（shopB||mj355 -> 收纳箱）",
      m.get("shopB||mj355") == "收纳箱", str(m.get("shopB||mj355")))
check("未收录仍回填兜底名", m.get("nomatch") == "謎の物体xyz", str(m.get("nomatch")))
check("空标题被跳过", "blank" not in m)
check("hit_count = 3", rm.get("hit_count") == 3, str(rm.get("hit_count")))
check("miss_count = 1", rm.get("miss_count") == 1, str(rm.get("miss_count")))
check("dict_total 与词库一致", rm.get("dict_total") == after, f"{rm.get('dict_total')} / {after}")
check("rematch 不写盘（条目数不变）",
      json.loads(urllib.request.urlopen(f"{API}/api/dictionary/stats", timeout=20).read())["total_count"] == after)

# 新导入的词条能立即被 rematch 命中（模拟「导入词库 -> 刷新品名」闭环）
suffix2 = uuid.uuid4().hex[:6]
newk = f"テスト圧泥器{suffix2}"
resp7 = post_multipart(f"{API}/api/dictionary/import", {},
                       {"file": ("patch.json", json.dumps({newk: "测试压泥器"}, ensure_ascii=False).encode("utf-8"))})
json.loads(resp7.read().decode("utf-8"))
resp8 = post_json(f"{API}/api/dictionary/rematch", {"items": [
    {"key": "x||1", "parent_sku": "1", "raw_name": f"2個 {newk} ステンレス製"}]})
rm2 = json.loads(resp8.read().decode("utf-8"))
check("导入新词后 rematch 立即可用", rm2["matched_names"].get("x||1") == "测试压泥器",
      str(rm2["matched_names"].get("x||1")))
check("数量词被跳过（新词生效）", rm2.get("miss_count") == 0, str(rm2.get("miss_count")))
urllib.request.urlopen(urllib.request.Request(
    f"{API}/api/dictionary?ja=" + urllib.parse.quote(newk), method="DELETE"), timeout=20).read()
check("清理测试词条",
      json.loads(urllib.request.urlopen(f"{API}/api/dictionary/stats", timeout=20).read())["total_count"] == after)

print("\n" + "=" * 78)
print(f"结果：PASS {OK}  FAIL {FAIL}")
print("=" * 78)
sys.exit(1 if FAIL else 0)
